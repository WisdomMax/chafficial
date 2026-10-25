#!/usr/bin/env python3
"""차피셜(CHA-FICIAL) DIRECTIVE 043: 딜러 실거래 Excel 입력 양식 생성 스크립트

장안평 제휴 딜러가 복잡한 서식 없이 엑셀로 간편하게 실거래 기록을 입력할 수 있도록
'거래기록 입력' 시트와 '작성 안내' 시트로 구성된 공식 템플릿을 생성한다.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def create_dealer_template(output_path: str = "data/dealer_pilot/차피셜_딜러_실거래_입력양식_v1.xlsx"):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = openpyxl.Workbook()

    # ----------------------------------------------------
    # Sheet 1: 거래기록 입력
    # ----------------------------------------------------
    ws1 = wb.active
    ws1.title = "거래기록 입력"
    ws1.views.sheetView[0].showGridLines = True

    # 스타일 정의
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # 진남색
    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    
    example_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # 연회색
    example_font = Font(name="맑은 고딕", size=10, italic=True, color="595959")
    
    data_font = Font(name="맑은 고딕", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = [
        "차량관리번호", "차종/트림", "연식", "주행거리(km)", 
        "매입일", "매입금액(원)", "입고일", "판매일", "판매금액(원)", 
        "취득비용(원)", "상품화비용(원)", "기타비용(원)", 
        "거래상태", "비용정산", "증빙확인", "데이터출처", "증거등급"
    ]

    # 헤더 작성
    for col_idx, h in enumerate(headers, start=1):
        cell = ws1.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws1.row_dimensions[1].height = 28

    # 예시 데이터 (2행)
    example_rows = [
        ["DTR_EX_01", "현대 아반떼 CN7 1.6 가솔린 모던", 2024, 35000, "2026-08-10", 15500000, "2026-08-12", "2026-09-05", 17800000, 450000, 350000, 150000, "판매완료", "완료", "가능", "DEALER_SELF_REPORTED", "D2"],
        ["DTR_EX_02", "현대 아반떼 CN7 1.6 인스퍼레이션", 2024, 18000, "2026-09-15", 19000000, "2026-09-17", "", "", 480000, 250000, "", "보유중", "미완료", "가능", "DEALER_SELF_REPORTED", "D1"]
    ]

    for r_idx, row_data in enumerate(example_rows, start=2):
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.fill = example_fill
            cell.font = example_font
            cell.border = thin_border
            if isinstance(val, (int, float)):
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center" if c_idx in [1, 3, 5, 7, 8, 13, 14, 15, 16, 17] else "left", vertical="center")
        ws1.row_dimensions[r_idx].height = 22

    # 실제 입력 빈 행 (10행 추가)
    for r_idx in range(4, 14):
        for c_idx in range(1, len(headers) + 1):
            cell = ws1.cell(row=r_idx, column=c_idx, value="")
            cell.font = data_font
            cell.border = thin_border
            if c_idx in [3, 4, 6, 9, 10, 11, 12]:
                cell.number_format = '#,##0'
        ws1.row_dimensions[r_idx].height = 22

    # 열 너비 자동 조정
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len * 1.5 + 3, 13)

    # ----------------------------------------------------
    # Sheet 2: 작성 안내
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="작성 안내")
    ws2.views.sheetView[0].showGridLines = True

    guide_title_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    title_font = Font(name="맑은 고딕", size=14, bold=True, color="1F4E79")
    sub_title_font = Font(name="맑은 고딕", size=11, bold=True, color="2F5597")
    body_font = Font(name="맑은 고딕", size=10)
    bullet_font = Font(name="맑은 고딕", size=10, bold=True)

    ws2.cell(row=1, column=1, value="차피셜(CHA-FICIAL) 제휴 딜러 실거래 기록 작성 가이드").font = title_font
    ws2.row_dimensions[1].height = 30

    guide_items = [
        ("1. 데이터 입력 기본 원칙", [
            "• 모든 금액은 '원 단위'로 입력해 주십시오 (예: 1,500만 원 -> 15000000).",
            "• 확인되지 않거나 모르는 금액/비용은 0원으로 적지 마시고, 반드시 '공란(빈칸)'으로 남겨주십시오.",
            "• 실제로 발생하지 않은 비용(0원 확정)인 경우에만 명시적으로 '0'을 입력합니다.",
            "• 현재 매장에 전시 및 보유 중인 차량은 판매일과 판매금액을 '공란'으로 비워둡니다.",
            "• 계약 체결 후 취소된 거래는 거래상태를 '취소'로 표기합니다."
        ]),
        ("2. 개인정보 및 민감정보 보호 원칙 (절대 기입 금지)", [
            "• 고객의 성명, 주민등록번호, 연락처(전화번호)는 절대로 기입하지 마십시오.",
            "• 실제 차량번호(예: 12가 3456) 및 차대번호(17자리)는 입력할 필요가 없으며 익명 관리번호만 사용합니다.",
            "• 제공해주신 데이터는 승인된 '차피셜 실거래 수익성 및 회전일수 분석' 목적으로만 안전하게 보관됩니다."
        ]),
        ("3. 증빙 자료 및 확인 등급 안내", [
            "• 원본 계약서나 금융 영수증 사본 첨부는 필수가 아니며, 상사 장부나 정산 메모 기반으로 작성 가능합니다.",
            "• 증빙확인 가능 여부에 따라 다음과 같이 등급이 관리됩니다:",
            "   - D1: 딜러 대표/담당자 구두 진술 및 상사 메모 수준",
            "   - D2: 매매계약서, 금융 이체내역, 세금계산서 중 1차 객관 증빙 확인",
            "   - D3: 매입가, 판매가, 상품화 영수증, 계약일 전수 대조 완료 (수익성 확정 가능)",
            "   - D4: 이전등록 및 취소/환불 기간 경과 세무 정산 완결"
        ]),
        ("4. 표본 수집 권장 범위 (생존 편향 방지)", [
            "• 수익이 크게 난 우량 매물만 골라서 제공하지 마시고, 최근 매입한 차량 전체(5~10대)를 있는 그대로 제공해 주십시오.",
            "• 빠른 판매 차량뿐만 아니라 현재 40일 이상 장기 보유 중인 차량, 상품화 비용이 예상보다 많이 들어 손실이 발생한 차량도 모델의 정확한 학습에 필수적인 귀중한 데이터입니다."
        ])
    ]

    cur_row = 3
    for sec_title, bullets in guide_items:
        t_cell = ws2.cell(row=cur_row, column=1, value=sec_title)
        t_cell.font = sub_title_font
        cur_row += 1
        for b in bullets:
            b_cell = ws2.cell(row=cur_row, column=1, value=b)
            b_cell.font = body_font
            cur_row += 1
        cur_row += 1

    ws2.column_dimensions['A'].width = 95

    wb.save(output_path)
    print(f"[+] 딜러용 공식 Excel 입력 양식 생성 완료: {output_path}")


if __name__ == "__main__":
    create_dealer_template()
